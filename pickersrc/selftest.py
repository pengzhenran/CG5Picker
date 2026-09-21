"""无界面自检 —— 解析 / 条件 / 列序 / 导出 / 理论潮汐 全链路跑一遍。

用法：
    python -m pickersrc --selftest <CG5手簿> [--out 输出目录]

给的是真手簿就能当回归测试用；不依赖 pytest、不弹窗。
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from . import cg5, exporter
from .colorder import ColOrder, default_visible
from .plots import solid_earth_tide_uGal
from .rules import (COND_KINDS, Cond, apply_conditions, default_conditions,
                    preview)

_OK = "  [OK]  "
_BAD = "  [FAIL]"


def _check(label: str, cond: bool, detail: str = "") -> bool:
    print(f"{_OK if cond else _BAD} {label}" + (f"　{detail}" if detail else ""))
    return cond


def run_selftest(argv: list[str]) -> int:
    files = [a for a in argv if not a.startswith("--") and a != "selftest"]
    out_dir = None
    if "--out" in argv:
        i = argv.index("--out")
        if i + 1 < len(argv):
            out_dir = Path(argv[i + 1])
            files = [f for f in files if f != argv[i + 1]]

    print("=" * 72)
    print("CG5 数据挑选 —— 无界面自检")
    print("=" * 72)

    ok = True

    # ---------------------------------------------------------------- 规范算例
    print("\n[1] 理论固体潮回归基准（DZ/T 0082-2021 附录 H.5）")
    val = solid_earth_tide_uGal(2003, 5, 6, 19.75, 31.0 + 20.0 / 60.0, 93.0)
    ok &= _check("规范算例 50.664 μGal", abs(val - 50.664) < 0.01,
                 f"算得 {val:.3f} μGal")

    # ---------------------------------------------------------------- 解析
    if not files:
        print("\n[2] 未给手簿文件 → 跳过解析 / 条件 / 导出 检查")
        print("\n提示：python -m pickersrc --selftest <CG5手簿.txt> [--out 目录]")
        return 0 if ok else 1

    for path in files:
        p = Path(path)
        print(f"\n[2] 解析  {p}")
        if not p.exists():
            ok &= _check("文件存在", False, str(p))
            continue
        f = cg5.read_file(p)
        ok &= _check("读到观测行", len(f.rows) > 0, f"{len(f.rows)} 条")
        ok &= _check("识别列头", bool(f.columns), " ".join(f.columns))
        ok &= _check("识别日期", bool(f.dates),
                     f"{len(f.dates)} 天，最新 {f.dates[0] if f.dates else '—'}")
        ok &= _check("日期按最新在前",
                     f.dates == sorted(f.dates, reverse=True) if f.dates else False)
        ok &= _check("每行都有日期", all(r.date for r in f.rows))
        ok &= _check("单元格数与列数一致",
                     all(len(r.cells) == len(f.columns) for r in f.rows))
        ok &= _check("日期与 DATE 列一致",
                     all(r.date == cg5.normalize_date(r.cells[-1]) for r in f.rows[:200]
                         if r.cells and r.cells[-1]))

        date = f.dates[0]
        svs = f.surveys_of(date)
        print(f"      最新日期 {date}：{f.count_of(date)} 条 / Survey {svs}")
        idx = f.subset(date, svs)
        ok &= _check("按日期+Survey 圈定", len(idx) == f.count_of(date),
                     f"{len(idx)} 条")

        # ------------------------------------------------------------ 条件
        print(f"\n[3] 条件筛选  {p.name} @ {date}")
        conds = default_conditions()
        print("      默认条件：" + "；".join(c.describe() for c in conds))
        res = preview(f, idx, conds)
        print(f"      {res['summary']}")
        for k, v in res["per_cond"].items():
            print(f"        命中 {v:>5}　{k}")
        ok &= _check("默认条件有结果", res["n_total"] > 0)
        ok &= _check("保留 + 剔除 = 总数",
                     res["n_keep"] + res["n_drop"] == res["n_total"])

        # 每条条件单独跑一遍，确保都能执行（不缺列、不报错）
        for kind, meta in COND_KINDS.items():
            c = Cond(kind, dict(meta.params))
            try:
                hits, _pc = apply_conditions(f, idx, [c])
                n_drop = sum(1 for h in hits if not h.keep)
                print(f"      单条件 {c.describe()}　剔除 {n_drop}")
            except Exception as exc:                    # noqa: BLE001
                ok &= _check(f"条件 {kind} 可执行", False, repr(exc))

        # 取消勾选 → 不参与筛选
        off = [Cond(c.kind, dict(c.params), enabled=False) for c in conds]
        r_off = preview(f, idx, off)
        ok &= _check("全部取消勾选后全部保留",
                     r_off["n_keep"] == r_off["n_total"], r_off["summary"])

        # ------------------------------------------------------------ 列序
        print(f"\n[4] 列顺序")
        vis = default_visible(f.columns)
        print("      默认列序：" + "、".join(vis))
        ok &= _check("默认列序以 线号/点号/读数/时长/时刻/日期 开头",
                     vis[:6] == ["线号", "点号", "读数(mGal)", "观测时长(s)",
                                 "时刻", "日期"], "、".join(vis[:6]))
        ok &= _check("列序不丢列", sorted(vis) == sorted(f.columns))

        # ------------------------------------------------------------ 导出
        print(f"\n[5] 导出 xlsx")
        keep_idx = [s for s, h in zip(idx, res["hits"]) if h.keep]
        dst = Path(out_dir) if out_dir else Path(tempfile.mkdtemp(prefix="cg5pick_"))
        dst.mkdir(parents=True, exist_ok=True)
        co = ColOrder(order=vis, reason=True)

        p1 = exporter.export_xlsx(f, keep_idx, co, dst / "原始格式.xlsx", mode="verbatim")
        p2 = exporter.export_xlsx(f, keep_idx, co, dst / "自定义列序.xlsx",
                                  mode="parsed", include_reason=True,
                                  reasons=[res["hits"][idx.index(s)].reason for s in keep_idx])
        for label, fp in (("原始格式", p1), ("自定义列序", p2)):
            n_row, n_col, head = _peek_xlsx(fp)
            ok &= _check(f"{label} 已写出", Path(fp).exists() and n_row > 0,
                         f"{Path(fp).name}：{n_row} 行 × {n_col} 列")
            ok &= _check(f"{label} 表头正确", bool(head), " ".join(head[:8]))
            ok &= _check(f"{label} 无前导空格",
                         all(not str(c).startswith(" ") for c in head))
        print(f"      输出目录：{dst}")

    print("\n" + "=" * 72)
    print("自检结果：" + ("全部通过 ✔" if ok else "存在失败项 ✘"))
    print("=" * 72)
    return 0 if ok else 1


def _peek_xlsx(path: str) -> tuple[int, int, list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    head = [str(x) for x in (next(rows, None) or [])]
    n = 0
    ncol = 0
    for r in rows:
        if r is None:
            continue
        n += 1
        ncol = max(ncol, len(r))
    wb.close()
    return n, ncol, head


if __name__ == "__main__":
    raise SystemExit(run_selftest(sys.argv[1:]))
